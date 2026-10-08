"""Portable Excel snapshots; external strings are always written as text."""

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from job_tracker.domain import ApplicationInput, Outcome, Stage
from job_tracker.services import Tracker

HEADERS = ["Company", "Role", "Applied on", "Stage", "Outcome", "Requisition ID", "Notes"]
FIELDS = ["company", "role", "applied_on", "stage", "outcome", "requisition_id", "notes"]


def sheet(workbook, name: str, headers: list[str], rows: list[list]):
    ws = workbook.create_sheet(name)
    ws.append(headers)
    for values in rows:
        ws.append(values)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Do not turn email text into spreadsheet formulas.
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="196E62")
    ws.freeze_panes = "A2"
    if rows:
        table = Table(displayName=name.replace(" ", "") + "Table", ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    for column in ws.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[column[0].column_letter].width = min(max(width + 3, 16), 55)


def export_search(tracker: Tracker, search_id: int, path: Path):
    searches = {row["id"]: row for row in tracker.searches()}
    if search_id not in searches:
        raise ValueError("Select a search to export.")
    apps = tracker.applications(search_id)
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet(workbook, "Applications", HEADERS, [[app[field] for field in FIELDS] for app in apps])
    dropdown(workbook["Applications"], "D", [item.value for item in Stage], "Stage")
    dropdown(workbook["Applications"], "E", [item.value for item in Outcome], "Outcome")
    events = [
        [app["company"], app["role"], item["kind"], item["effective_at"], item["evidence"]]
        for app in apps
        for item in tracker.events(app["id"])
    ]
    sheet(workbook, "Events", ["Company", "Role", "Event", "Date UTC", "Evidence"], events)
    tasks = [
        [task["company"], task["role"], task["title"], task["due_at"], task["completed"]]
        for task in tracker.tasks(search_id)
    ]
    sheet(workbook, "Tasks", ["Company", "Role", "Task", "Due UTC", "Completed"], tasks)
    dropdown(workbook["Tasks"], "E", ["TRUE", "FALSE"], "Completed")
    sheet(
        workbook,
        "Search",
        ["Name", "Start date", "End date"],
        [[searches[search_id][key] for key in ("name", "start_date", "end_date")]],
    )
    workbook.save(path)
    workbook.close()


def dropdown(ws, column: str, options: list[str], label: str):
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(options) + '"',
        allow_blank=False,
        showDropDown=False,
        showErrorMessage=True,
        showInputMessage=True,
    )
    validation.errorTitle = f"Choose a valid {label.lower()}"
    validation.error = "Select a value from the dropdown."
    validation.promptTitle = label
    validation.prompt = "Choose from the list. Excel edits do not automatically update the app."
    validation.errorStyle = "stop"
    ws.add_data_validation(validation)
    validation.add(f"{column}2:{column}1048576")


def preview_import(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = workbook["Applications"] if "Applications" in workbook.sheetnames else workbook.active
        iterator = ws.iter_rows(values_only=True)
        header = next(iterator, ())
        headers = [str(value or f"Column {i + 1}") for i, value in enumerate(header)]
        rows = [list(row) for row in iterator if any(value is not None for value in row)]
        if len(rows) > 10000:
            raise ValueError("Import at most 10,000 rows at a time.")
        samples = [[str(value or "") for value in row] for row in rows[:3]]
        return {"headers": headers, "rows": rows, "samples": samples, "count": len(rows)}
    finally:
        workbook.close()


def import_rows(tracker: Tracker, search_id: int, preview: dict, mapping: dict) -> tuple[int, int]:
    if mapping.get("company", -1) < 0 or mapping.get("role", -1) < 0:
        raise ValueError("Map the company and role columns.")
    candidates = []
    for number, row in enumerate(preview["rows"], start=2):
        values = {}
        for field, index in mapping.items():
            value = row[index] if 0 <= index < len(row) else ""
            if isinstance(value, (date, datetime)):
                value = value.strftime("%Y-%m-%d")
            values[field] = str(value).strip() if value is not None else ""
        values["stage"] = values.get("stage") or Stage.APPLIED
        values["outcome"] = values.get("outcome") or Outcome.ACTIVE
        # Validate all rows before modifying the database.
        try:
            candidates.append(ApplicationInput.model_validate(values))
        except ValueError as error:
            raise ValueError(f"Row {number} is invalid; check dates and status columns.") from error
    known = {
        (
            app["company"].casefold(),
            app["role"].casefold(),
            app["applied_on"],
            app["requisition_id"].casefold(),
        )
        for app in tracker.applications(search_id)
    }
    added = skipped = 0
    # Entire import is atomic. Do not silently overwrite existing manual records.
    from job_tracker.db import Application, ApplicationEvent
    from job_tracker.domain import now

    with tracker.db.sessions.begin() as session:
        for item in candidates:
            identity = (
                item.company.casefold(),
                item.role.casefold(),
                item.applied_on,
                item.requisition_id.casefold(),
            )
            if identity in known:
                skipped += 1
                continue
            app = Application(search_id=search_id, **item.model_dump(mode="json"))
            session.add(app)
            session.flush()
            session.add(
                ApplicationEvent(
                    application_id=app.id,
                    kind="manual",
                    evidence="Imported from Excel.",
                    effective_at=now(),
                )
            )
            known.add(identity)
            added += 1
    return added, skipped
